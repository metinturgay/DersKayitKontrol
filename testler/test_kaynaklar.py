# -*- coding: utf-8 -*-
"""Belge toplama akışı testleri.

Neyi koruyor
------------
Bu akış danışmanın kendi belgelerini verdiği yer. Sessiz kalırsa en
pahalı hatayı üretir: yanlış belge ya da yanlış YIL kabul edilirse
pano dolar, sayılar makul görünür ve hepsi yanlıştır.

Dört değişmez sınanıyor:

  1. Belge ALINIR ALINMAZ çözümlenir; anlaşılmayan belge KABUL EDİLMEZ.
  2. Giriş yılı asla sessizce tahmin edilmez - teyit ettirilir.
  3. DAMGA EN SON atılır: kurulum yarıda kalırsa (bozuk profil, yazma
     hatası) yerel kurulum devreye GİRMEZ.
  4. Etkileşimsiz çalışmada (script'ten çağırma) akış ASILI KALMAZ.
"""
import io
import json
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

BURASI = os.path.dirname(os.path.abspath(__file__))
KOK = os.path.dirname(BURASI)
sys.path.insert(0, KOK)
sys.path.insert(0, BURASI)

import kaynaklar                                          # noqa: E402
import sentetik_mufredat as sm                            # noqa: E402
import yollar                                             # noqa: E402

OK = []

# Bölümün BELGELERİ yayımlanan ağaçta yok (yalnız çözülmüş hâlleri
# gidiyor). Varsa gerçeğini, yoksa sentetiğini kullanıyoruz; test her
# iki durumda da aynı değişmezleri sınıyor.
DOCX = None
XLSX = None
SENTETIK = False


def _docx_hazirla(klasor):
    global DOCX, SENTETIK
    gercek = os.path.join(KOK, "veri", "mufredat.docx")
    if os.path.exists(gercek):
        DOCX = gercek
        return
    SENTETIK = True
    yariyillar = {}
    for yy in range(1, 9):
        yariyillar[yy] = [
            {"kod": "990%d10%d" % (yy, i), "ad": "Ders %d-%d" % (yy, i),
             "akts": 6} for i in range(1, 6)]
    DOCX = sm.belge_yaz(os.path.join(klasor, "sentetik.docx"), yariyillar)


def _xlsx_hazirla(klasor):
    global XLSX
    gercek = os.path.join(KOK, "veri", "ders_programi.xlsx")
    if os.path.exists(gercek):
        XLSX = gercek
        return
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "I. Öğretim"
    saatler = ("08:30-09:15", "09:25-10:10", "10:20-11:05", "11:15-12:00")
    for i, saat in enumerate(saatler):
        ws.cell(1, 3 + i).value = saat
    ws.cell(2, 1).value = "PAZARTESİ"
    ws.cell(2, 2).value = "1"
    ws.cell(2, 3).value = "Ders 1-1 (M1)"
    ws.cell(2, 4).value = "Ders 1-1 (M1)"
    ws.cell(3, 2).value = "2"
    ws.cell(3, 3).value = "Ders 3-1 (M2)"
    ws.cell(3, 4).value = "Ders 3-2 (M2)"
    ws.cell(4, 1).value = "SALI"
    ws.cell(4, 2).value = "1"
    ws.cell(4, 5).value = "Ders 1-2 (M1)"
    XLSX = os.path.join(klasor, "sentetik.xlsx")
    wb.save(XLSX)


def kontrol(ad, beklenen, gelen):
    tamam = beklenen == gelen
    print(("  [OK]  " if tamam else "  [HATA] ") + ad +
          ("" if tamam else "\n          beklenen=%r\n          gelen   =%r"
           % (beklenen, gelen)))
    OK.append(tamam)


def _cevaplayici(cevaplar):
    """Scripted input: her çağrıda sıradaki cevabı verir."""
    akis = iter(cevaplar)

    def sor(istem):
        return next(akis, "tamam")
    return sor


def _bozuk_docx(klasor):
    """Tanınmayan şablonlu bir docx (ders çıkmaz)."""
    g = ('<w:p><w:r><w:t>Ders Listesi</w:t></w:r></w:p>'
         '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Sütun1</w:t></w:r></w:p>'
         '</w:tc></w:tr></w:tbl>')
    x = ('<?xml version="1.0"?><w:document xmlns:w="%s"><w:body>%s'
         '</w:body></w:document>' % (sm.W, g))
    p = os.path.join(klasor, "bozuk.docx")
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("[Content_Types].xml", sm._TIPLER)
        z.writestr("_rels/.rels", sm._ILISKILER)
        z.writestr("word/document.xml", x.encode("utf-8"))
    return p


def _yol_temizleme():
    print("")
    print("  === Yol temizleme ===")
    # Dosyayı pencereye sürüklemek yolu TIRNAK İÇİNDE veriyor; bu
    # temizlenmezse "bulunamadı" denip kullanıcı haklı olarak şaşırıyor.
    kontrol("çift tırnak soyuluyor", "C:/a b/c.docx",
            kaynaklar.yol_temizle('"C:/a b/c.docx"'))
    kontrol("tek tırnak soyuluyor", "C:/a/c.docx",
            kaynaklar.yol_temizle("'C:/a/c.docx'"))
    kontrol("baş/son boşluk gidiyor", "C:/a.docx",
            kaynaklar.yol_temizle("  C:/a.docx  "))
    kontrol("içerideki boşluk KORUNUYOR", "C:/a b/c d.docx",
            kaynaklar.yol_temizle('  "C:/a b/c d.docx" '))
    kontrol("boş giriş", "", kaynaklar.yol_temizle("   "))


def _belge_dogrulama(gecici):
    print("")
    print("  === Belge doğrulama ===")
    tamam, ozet, veri = kaynaklar.belge_dogrula(DOCX)
    kontrol("geçerli belge kabul ediliyor", True, tamam)
    kontrol("ne okunduğu yazılıyor (ders + AKTS)", True,
            " ders " in ozet and "AKTS" in ozet and "yarıyıl" in ozet)

    tamam, mesaj, _ = kaynaklar.belge_dogrula(
        os.path.join(gecici, "yok.docx"))
    kontrol("olmayan dosya reddediliyor", False, tamam)
    kontrol("sebebi söyleniyor", True, "Bulunamadı" in mesaj)

    tamam, mesaj, _ = kaynaklar.belge_dogrula(XLSX)
    kontrol("yanlış uzantı reddediliyor", False, tamam)
    kontrol("hangi uzantı beklendiği yazılıyor", True, ".docx" in mesaj)

    # ASIL sınama: okunabilen ama ANLAŞILMAYAN belge.
    bozuk = _bozuk_docx(gecici)
    tamam, mesaj, _ = kaynaklar.belge_dogrula(bozuk)
    kontrol("anlaşılmayan belge reddediliyor", False, tamam)
    kontrol("hata belgeden ne görüldüğünü yazıyor", True,
            "ders" in mesaj.lower())


def _program_dogrulama(gecici):
    print("")
    print("  === Ders programı doğrulama ===")
    tamam, ozet, _ = kaynaklar.program_dogrula(XLSX)
    kontrol("geçerli program kabul ediliyor", True, tamam)
    kontrol("ne okunduğu yazılıyor (ders + çift)", True,
            " ders " in ozet and "çift" in ozet)

    tamam, mesaj, _ = kaynaklar.program_dogrula(DOCX)
    kontrol("yanlış uzantı reddediliyor", False, tamam)

    import openpyxl
    bos = os.path.join(gecici, "bos.xlsx")
    openpyxl.Workbook().save(bos)
    tamam, mesaj, _ = kaynaklar.program_dogrula(bos)
    kontrol("bomboş program reddediliyor", False, tamam)
    kontrol("sebebi söyleniyor", True,
            "saat" in mesaj.lower() or "ders" in mesaj.lower())

    # AYRI bir durum: yerleşim BULUNUYOR (saat başlıkları ve gün adı
    # var) ama hücrelerde hiç ders yok. Bunu yukarıdaki "saat sütunu
    # yok" kapısı yakalamaz; "hiç ders okunamadı" kapısı yakalamalı.
    # Mutasyon denemesinde bu kapıyı kaldırdığımda hiçbir test kızmadı.
    iskelet = os.path.join(gecici, "iskelet.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "I. Öğretim"
    for i, saat in enumerate(("08:30-09:15", "09:25-10:10", "10:20-11:05",
                              "11:15-12:00")):
        ws.cell(1, 3 + i).value = saat
    ws.cell(2, 1).value = "PAZARTESİ"
    ws.cell(2, 2).value = "1"
    ws.cell(3, 2).value = "2"
    wb.save(iskelet)
    tamam, mesaj, _ = kaynaklar.program_dogrula(iskelet)
    kontrol("dersi olmayan program reddediliyor", False, tamam)
    kontrol("hiç ders okunamadığı söyleniyor", True,
            "ders" in mesaj.lower())


def _yil_sorusu():
    print("")
    print("  === Giriş yılı ASLA tahmin edilmiyor ===")
    # Dosya adında yıl VAR: teyit ettirilir.
    kontrol("dosya adındaki yıl onaylanınca kabul", 2024,
            kaynaklar.yil_sor("2024-2025 Okutulacak Dersler.docx",
                              _cevaplayici(["e"])))
    # Onaylanmazsa elle sorulur.
    kontrol("onaylanmazsa elle sorulur", 2021,
            kaynaklar.yil_sor("2024-2025 Okutulacak Dersler.docx",
                              _cevaplayici(["h", "2021"])))
    # Dosya adında yıl YOK: doğrudan sorulur, saçma cevap kabul edilmez.
    kontrol("yılsız dosyada saçma cevap reddedilir", 2023,
            kaynaklar.yil_sor("mufredat.docx",
                              _cevaplayici(["abc", "12", "2023"])))


def _belge_dongusu(gecici):
    print("")
    print("  === Belge döngüsü ===")
    bozuk = _bozuk_docx(gecici)

    # Tek belge, dosya adında yıl yok -> elle
    secilen = kaynaklar.belge_dongusu(
        _cevaplayici([DOCX, "2026", "tamam"]))
    kontrol("tek belge toplandı", {2026: DOCX}, secilen)

    # Tırnak içinde verilen yol da çalışmalı
    secilen = kaynaklar.belge_dongusu(
        _cevaplayici(['"%s"' % DOCX, "2025", "tamam"]))
    kontrol("tırnaklı yol kabul ediliyor", {2025: DOCX}, secilen)

    # Bozuk belge REDDEDİLİR ama döngü sürer
    secilen = kaynaklar.belge_dongusu(
        _cevaplayici([bozuk, DOCX, "2024", "tamam"]))
    kontrol("bozuk belge alınmadı, döngü sürdü",
            {2024: DOCX}, secilen)

    # Hiç belge vermeden 'tamam': kabul EDİLMEZ
    secilen = kaynaklar.belge_dongusu(
        _cevaplayici(["tamam", DOCX, "2023", "tamam"]))
    kontrol("belgesiz bitirilemiyor", {2023: DOCX}, secilen)

    # Aynı yıl iki kez: sonuncusu geçerli
    secilen = kaynaklar.belge_dongusu(
        _cevaplayici([DOCX, "2022", DOCX, "2022", "tamam"]))
    kontrol("aynı yıl tekrarında tek kayıt", 1, len(secilen))


def _program_sorusu(gecici):
    print("")
    print("  === Ders programı sorusu ===")
    kontrol("boş bırakılabiliyor", "",
            kaynaklar.program_sor(_cevaplayici([""])))
    kontrol("geçerli dosya kabul", XLSX,
            kaynaklar.program_sor(_cevaplayici([XLSX])))
    # Kötü dosya reddedilir, tekrar sorulur
    kontrol("kötü dosyadan sonra tekrar soruluyor", XLSX,
            kaynaklar.program_sor(
                _cevaplayici([DOCX, XLSX])))


CEVAPLAR = {
    "universite": "Selçuk Üniversitesi", "fakulte": "Sınama Fakültesi",
    "ad": "Sınama", "program_yili": 4, "ogretim_yili": "2026-2027",
    "tos_azami_adet": 1, "acik_kapatma_yariyili": 5,
    "asenkron_dersler": [], "esnek_dersler": [], "uygulamali_dersler": [],
}


def _kurulum():
    print("")
    print("  === Kurulum: DAMGA EN SON ===")
    kok = tempfile.mkdtemp(prefix="dkk_kur_")
    eski = yollar.yazilan_kok
    try:
        yollar.yazilan_kok = lambda: Path(kok)
        yollar._damga_onbellek = None

        # 1. BOZUK profil: kurulum başarısız, DAMGA ATILMAMALI
        kotu = dict(CEVAPLAR, program_yili=3)      # 2/4/5 değil
        tamam, mesaj = kaynaklar.kur({2026: DOCX}, "", kotu,
                                     yaz=lambda *a: None)
        kontrol("bozuk profille kurulum başarısız", False, tamam)
        kontrol("sebebi söyleniyor", True, "program_yili" in mesaj)
        damga = Path(kok) / yollar.YEREL_KLASOR / yollar.DAMGA
        kontrol("BAŞARISIZ kurulumda damga YOK", False, damga.exists())
        yollar._damga_onbellek = None
        kontrol("yarım kurulum devreye girmiyor", False,
                yollar.yerel_hazir_mi())

        # 2. Geçerli kurulum
        tamam, mesaj = kaynaklar.kur({2026: DOCX}, XLSX,
                                     CEVAPLAR, yaz=lambda *a: None)
        kontrol("geçerli kurulum başarılı", True, tamam)
        yollar._damga_onbellek = None
        kontrol("damga atıldı", True, damga.exists())
        kontrol("yerel kurulum hazır", True, yollar.yerel_hazir_mi())

        d = json.loads(damga.read_text(encoding="utf-8"))
        kontrol("damga bölümü yazıyor", "Sınama", d.get("bolum"))
        kontrol("damga belgeleri yazıyor", ["2026"],
                sorted((d.get("belgeler") or {}).keys()))
        # Arşiv KLASÖRÜ de sahiplenilmeli; yoksa belgeler gömülü kökte
        # aranır ve türetilenler hiç üretilmez (ölçüldü).
        for ad in ("bolum.json", "mufredat", "mufredat.json",
                   "cakismalar.json", "ders_programi.xlsx"):
            kontrol("damga '%s' sahipleniyor" % ad, True,
                    ad in (d.get("dosyalar") or []))

        yerel = Path(kok) / yollar.YEREL_KLASOR
        kontrol("belge kopyalandı", True,
                (yerel / "mufredat" / "2026.docx").exists())
        kontrol("program kopyalandı", True,
                (yerel / "ders_programi.xlsx").exists())
        kontrol("profil yazıldı", True, (yerel / "bolum.json").exists())
        # Çözülmüş hâller yeniden başlatmadan SONRA üretilir; kurulum
        # onları YAZMAMALI (eski profille üretilmiş olurlardı).
        kontrol("çözülmüş dosyalar kurulumda üretilmiyor", False,
                (yerel / "mufredat.json").exists())
    finally:
        yollar.yazilan_kok = eski
        yollar._damga_onbellek = None
        shutil.rmtree(kok, ignore_errors=True)


def _etkilesimsiz():
    print("")
    print("  === Etkileşimsiz çalışma ===")
    # stdin bir terminale bağlı değilse soru SORULMAZ: script'ten
    # çağrılan exe input() üzerinde asılı kalmamalı.
    kaynak = io.open(os.path.join(KOK, "kaynaklar.py"),
                     encoding="utf-8").read()
    kontrol("etkilesim_var stdin'e bakıyor", True,
            "isatty()" in kaynak)
    b = io.open(os.path.join(KOK, "baslat.py"), encoding="utf-8").read()
    kontrol("teyit etkileşimsizde atlanıyor", True,
            "etkilesim_var()" in b)
    # Fonksiyonun KENDİSİNİ sınıyoruz; bu oturumun stdin'i neyse onu
    # değil. (İlk yazışta "bu testte stdin terminal değil" diye
    # yazılmıştı ve kabuktan çalışınca haklı olarak kırıldı.)
    class _Sahte(object):
        def __init__(self, deger):
            self._deger = deger

        def isatty(self):
            if isinstance(self._deger, Exception):
                raise self._deger
            return self._deger

    asil = sys.stdin
    try:
        sys.stdin = _Sahte(False)
        kontrol("terminal değilse soru sorulmaz", False,
                kaynaklar.etkilesim_var())
        sys.stdin = _Sahte(True)
        kontrol("terminalse soru sorulur", True,
                kaynaklar.etkilesim_var())
        sys.stdin = None
        kontrol("stdin yoksa soru sorulmaz", False,
                kaynaklar.etkilesim_var())
        sys.stdin = _Sahte(ValueError("kapali"))
        kontrol("stdin patlarsa soru sorulmaz", False,
                kaynaklar.etkilesim_var())
    finally:
        sys.stdin = asil


def main():
    gecici = tempfile.mkdtemp(prefix="dkk_kaynak_")
    try:
        _docx_hazirla(gecici)
        _xlsx_hazirla(gecici)
        if SENTETIK:
            print("  (bölüm belgeleri yok; sentetik belgelerle "
                  "çalışılıyor)")
        _yol_temizleme()
        _belge_dogrulama(gecici)
        _program_dogrulama(gecici)
        _yil_sorusu()
        _belge_dongusu(gecici)
        _program_sorusu(gecici)
        _kurulum()
        _etkilesimsiz()
    finally:
        shutil.rmtree(gecici, ignore_errors=True)
    print("")
    print("  %d/%d kontrol geçti." % (sum(OK), len(OK)))
    return 0 if all(OK) else 1


if __name__ == "__main__":
    sys.exit(main())
