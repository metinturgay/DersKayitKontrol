# -*- coding: utf-8 -*-
"""Ders programı okumasının denetimi.

Excel'den hiçbir şey kaçırmadığımızı doğrular:
  - birleşik (merged) hücreler yüzünden atlanan saat var mı
  - dolu olduğu hâlde ders olarak okunmayan hücreler
  - programda olup katalogda olmayan / katalogda olup programda olmayan
  - aynı saatte aynı derslikte iki ders (program hatası)

    python program_denetle.py
"""
import glob
from pathlib import Path

import ders_programi as dp

BURASI = Path(__file__).resolve().parent


def _xlsx_bul():
    for kalip in (str(BURASI / "veri" / "*.xlsx"),
                  str(Path.home() / "Desktop" / "*Ders Program*.xlsx")):
        adaylar = sorted(glob.glob(kalip))
        if adaylar:
            return adaylar[0]
    raise SystemExit("Ders programı xlsx bulunamadı.")


def main():
    import openpyxl

    yol = _xlsx_bul()
    wb = openpyxl.load_workbook(yol, data_only=True)
    ws = wb["I. Öğretim"]
    print("Dosya : " + yol)
    print("Sayfa : I. Öğretim  (%d satır x %d sütun)"
          % (ws.max_row, ws.max_column))
    print("")

    sorunlar = []

    # ------------------------------------------------------------------
    # 1. Birleşik hücreler
    # ------------------------------------------------------------------
    # Kritik ayrım:
    #   DİKEY birleşme (tek sütun, çok satır) -> aynı saat, farklı paralel
    #     ders satırları. Saat sütunu değişmediği için VERİ KAYBI YOK.
    #   YATAY birleşme (çok sütun) -> ders birden çok saate yayılıyor ama
    #     openpyxl yalnızca sol üstü okur, kalan SAATLER KAÇAR.
    saat_sutunlari = {ord(h) - 64 for h in dp.SAAT_SUTUNLARI}
    dikey, yatay = [], []
    for aralik in ws.merged_cells.ranges:
        sutunlar = set(range(aralik.min_col, aralik.max_col + 1))
        if not (sutunlar & saat_sutunlari):
            continue                       # gün/sınıf/dipnot alanı
        if aralik.min_col < 3:
            continue    # A/B sütunundan başlayanlar başlık, dipnot, imza
        deger = ws.cell(aralik.min_row, aralik.min_col).value
        if not deger:
            continue
        ad, _, _ = dp._ders_adi(deger)
        if not ad or ad.upper() in dp.YOKSAY:
            continue                       # imza, dipnot vb.
        kayit = (str(aralik), str(deger).split("\n")[0][:38])
        if len(sutunlar & saat_sutunlari) > 1:
            yatay.append(kayit + (sorted(sutunlar & saat_sutunlari),))
        else:
            dikey.append(kayit)

    print("1. BİRLEŞİK HÜCRELER")
    print("   Toplam birleşik alan          : %d" % len(ws.merged_cells.ranges))
    print("   Ders taşıyan, DİKEY (zararsız): %d" % len(dikey))
    print("   Ders taşıyan, YATAY (riskli)  : %d" % len(yatay))
    print("")
    print("   Dikey birleşme tek sütunda kalır; saat sütunu değişmediği")
    print("   için okunan saatler eksilmez. Yatay birleşme ise dersi birden")
    print("   çok saate yayar ve openpyxl yalnızca ilk hücreyi okur.")
    if yatay:
        print("")
        print("   KAYIP RİSKİ OLAN YATAY BİRLEŞMELER:")
        for a, d, sut in yatay:
            print("     %-12s %-38s sütunlar=%s" % (a, d, sut))
        sorunlar.append("%d yatay birleşme saat kaçırıyor" % len(yatay))
    else:
        print("")
        print("   Yatay birleşme yok -> hiçbir saat kaçmıyor.")
        if dikey:
            print("   (Dikey örnekler: %s)"
                  % ", ".join("%s %s" % (a, d) for a, d in dikey[:3]))
    print("")

    # ------------------------------------------------------------------
    # 2. Okunamayan dolu hücreler
    # ------------------------------------------------------------------
    print("2. OKUNAMAYAN HÜCRELER")
    okunamayan = []
    saat_degerleri = set(dp.SAAT_SUTUNLARI.values())
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ham = ws.cell(r, c).value
            if ham is None or not str(ham).strip():
                continue
            metin = str(ham).strip()
            harf = chr(64 + c)
            if harf not in dp.SAAT_SUTUNLARI:
                continue                        # gün/sınıf/dipnot sütunları
            if metin in saat_degerleri:
                continue                        # saat başlığı
            ad, _, _ = dp._ders_adi(ham)
            if ad and ad.upper() in dp.YOKSAY:
                continue                    # bilerek atlanan (imza vb.)
            if not ad:
                okunamayan.append((r, harf, metin.split("\n")[0][:50]))
    if okunamayan:
        print("   %d hücre ders olarak okunamadı:" % len(okunamayan))
        for r, h, m in okunamayan[:15]:
            print("     r%-3d %s: %s" % (r, h, m))
        sorunlar.append("%d okunamayan hücre" % len(okunamayan))
    else:
        print("   Saat sütunlarındaki dolu hücrelerin hepsi okundu.")
        print("   (Bilerek atlananlar: %s)" % ", ".join(sorted(dp.YOKSAY - {""})))
    print("")

    # ------------------------------------------------------------------
    # 3. Program <-> katalog karşılaştırması
    # ------------------------------------------------------------------
    print("3. PROGRAM <-> KATALOG")
    veri = dp.yukle()
    if not veri:
        print("   veri/cakismalar.json yok, atlandı.")
    else:
        katalog = _katalog_oku()
        if katalog is None:
            print("   cikti/ altında ders sayfası yok, atlandı.")
        else:
            harita = dp.program_eslestir(katalog, veri)
            eslesmeyen_katalog = [d for d in katalog
                                  if d["ders_no"] not in harita]
            eslesen_program = {h["program_adi"] for h in harita.values()}
            eslesmeyen_program = [a for a in veri["dersler"]
                                  if a not in eslesen_program]

            print("   Katalogdaki ders          : %d" % len(katalog))
            print("   Programdaki ders          : %d" % len(veri["dersler"]))
            print("   Eşleşen                   : %d" % len(harita))

            if eslesmeyen_katalog:
                print("")
                print("   Katalogda VAR, programda YOK (%d):"
                      % len(eslesmeyen_katalog))
                for d in eslesmeyen_katalog:
                    print("     %s %-42s %s" % (d["ders_no"],
                                                d["ders_adi"][:42],
                                                d.get("donem") or ""))
                sorunlar.append("%d katalog dersi programda yok"
                                % len(eslesmeyen_katalog))
            if eslesmeyen_program:
                print("")
                print("   Programda VAR, katalogda YOK (%d):"
                      % len(eslesmeyen_program))
                for a in eslesmeyen_program:
                    d = veri["dersler"][a]
                    print("     %-38s sınıf=%s %s" % (
                        a[:38], d["siniflar"],
                        "(asenkron)" if d.get("asenkron") else
                        "(esnek)" if d.get("esnek") else ""))
            # Benzerlikle (kesin olmayan) eşleşenler danışman görsün
            benzer = [(k, h) for k, h in harita.items() if h["tip"] == "benzer"]
            if benzer:
                print("")
                print("   Ad birebir aynı DEĞİL, benzerlikle eşleşti (%d)"
                      " — doğrulayın:" % len(benzer))
                ad_dizini = {d["ders_no"]: d["ders_adi"] for d in katalog}
                for k, h in benzer:
                    print("     %%%d  %-36s -> %s"
                          % (round(h["benzerlik"] * 100),
                             ad_dizini[k][:36], h["program_adi"]))
    print("")

    # ------------------------------------------------------------------
    # 4. Aynı saatte aynı derslik
    # ------------------------------------------------------------------
    print("4. DERSLİK ÇAKIŞMASI (program hatası göstergesi)")
    # Dersliği ders bazında değil, HÜCRE bazında izliyoruz: bir ders
    # farklı günlerde farklı derslikte olabiliyor, ders bazında bakmak
    # yanlış alarm üretiyor.
    yer = {}
    gun = None
    for r in range(1, ws.max_row + 1):
        a = dp._gun_adi(ws.cell(r, 1).value)
        if a:
            for g in dp.GUNLER:
                if a.startswith(g[:4]):
                    gun = g
                    break
        if gun is None:
            continue
        for c in range(3, 14):
            saat = dp.SAAT_SUTUNLARI.get(chr(64 + c))
            if not saat:
                continue
            ad, derslik, grup = dp._ders_adi(ws.cell(r, c).value)
            if not ad or ad.upper() in dp.YOKSAY or not derslik:
                continue
            yer.setdefault((gun, saat, derslik), set()).add(ad)
    catisan = {k: v for k, v in yer.items() if len(v) > 1}
    if catisan:
        print("   %d yer/saat birden çok ders taşıyor:" % len(catisan))
        for (g, s, d), adlar in sorted(catisan.items())[:10]:
            print("     %-10s %-12s %-5s : %s"
                  % (g, s, d, ", ".join(sorted(adlar))))
        sorunlar.append("%d derslik çakışması" % len(catisan))
    else:
        print("   Aynı derslikte aynı saatte iki ders yok.")
    print("")

    # ------------------------------------------------------------------
    print("=" * 62)
    if sorunlar:
        print("DİKKAT EDİLECEKLER: " + "; ".join(sorunlar))
    else:
        print("Okuma temiz görünüyor, veri kaybı tespit edilmedi.")
    return 0


def _katalog_oku():
    """TÜM öğrencilerin katalogunun birleşimi.

    Katalog öğrenciye göre değişiyor (herkes her dersi ekleyemiyor);
    tek öğrenciye bakmak "programda var, katalogda yok" listesini
    şişiriyor.
    """
    import ders_kayit as dk
    sayfalar = sorted((BURASI / "cikti").glob("ders_sayfasi_*.html"))
    if not sayfalar:
        return None
    birlesik = {}
    for yol in sayfalar:
        kayit = dk.ders_kaydini_coz(yol.read_text(encoding="utf-8"))
        for d in kayit["katalog"]:
            if d.get("sekme") == "Bölüm Dersleri":
                birlesik.setdefault(d["ders_no"], d)
    return sorted(birlesik.values(), key=lambda d: d["ders_no"])


if __name__ == "__main__":
    raise SystemExit(main())
